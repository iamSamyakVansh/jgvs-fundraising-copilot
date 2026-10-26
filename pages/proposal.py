"""Proposal editor, live budget, and approval submission."""
from io import BytesIO

import streamlit as st

from models import ApprovalItem
from services.budget import line_percentages, summarize_budget
from services.drafts import concept_note
from services.export import escape_html
from state import add_approval, add_proposal_version, get_proposal_versions
from ui.components import page_header


def _docx_bytes(markdown: str) -> bytes:
    """Create a simple Word document from line-oriented Markdown text."""
    from docx import Document
    document = Document()
    for line in markdown.splitlines():
        if line.startswith("# "): document.add_heading(line[2:], level=1)
        elif line.startswith("## "): document.add_heading(line[3:], level=2)
        elif line.startswith("- "): document.add_paragraph(line[2:], style="List Bullet")
        else: document.add_paragraph(line.replace("**", ""))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render() -> None:
    """Render editable proposal fields and submit immutable draft versions for review."""
    page_header("Proposal & Budget", "Shape a concept note with a transparent budget; unknowns remain clearly marked.")
    left, right = st.columns([1.2, 1])
    with left:
        project = st.text_input("Project title", "Community Digital Learning Initiative")
        geography = st.text_input("Project geography", "Information Required")
        duration = st.number_input("Duration (months)", 1, 60, 12)
        beneficiaries = st.text_input("Target beneficiaries", "Information Required")
        objective = st.text_area("Project objective", "Describe the intended change using verified JGVS project information.", height=90)
        evidence = st.text_area("Need and evidence", "Information Required", height=72)
        activities = st.text_area("Proposed activities", "Information Required", height=72)
        outcomes = st.text_area("Expected outcomes", "Information Required", height=72)
    with right:
        st.subheader("Budget (INR)")
        requested = st.number_input("Requested amount", min_value=0, value=1_000_000, step=50_000)
        items = {
            "Personnel": st.number_input("Personnel", min_value=0, value=300_000, step=10_000),
            "Equipment & materials": st.number_input("Equipment & materials", min_value=0, value=250_000, step=10_000),
            "Training & implementation": st.number_input("Training & implementation", min_value=0, value=250_000, step=10_000),
            "Travel & field work": st.number_input("Travel & field work", min_value=0, value=100_000, step=10_000),
            "Monitoring & administration": st.number_input("Monitoring & administration", min_value=0, value=100_000, step=10_000),
        }
        summary = summarize_budget(items, requested)
        st.metric("Line-item subtotal", f"₹{summary.subtotal:,.0f}", delta=f"Variance ₹{summary.variance:,.0f}")
        for label, percent in line_percentages(items).items():
            st.markdown(f"**{escape_html(label)}** · {percent:.1f}%")
            st.progress(percent / 100)
        if summary.variance:
            st.warning("Line items do not match the requested total. Reconcile the difference before approval.")
    if st.button("Generate and send to human review", type="primary"):
        if not project.strip():
            st.error("Enter a project title.")
        else:
            markdown = concept_note({"project": project.strip(), "geography": geography, "duration": duration, "beneficiaries": beneficiaries, "objective": objective, "evidence": evidence, "activities": activities, "outcomes": outcomes, "items": items, "requested": requested})
            versions = get_proposal_versions()
            version = len(versions) + 1
            add_proposal_version({"version": version, "title": project.strip(), "content": markdown})
            add_approval(ApprovalItem("Proposal draft", project.strip(), markdown, version=version))
            st.session_state["latest_proposal"] = markdown
            st.toast("Draft added to the approval queue.")
    draft = st.session_state.get("latest_proposal")
    if draft:
        st.subheader("Draft preview")
        st.markdown(draft)
        c1, c2 = st.columns(2)
        c1.download_button("Download Markdown", draft, "jgvs_concept_note.md", "text/markdown")
        try:
            c2.download_button("Download DOCX", _docx_bytes(draft), "jgvs_concept_note.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        except ImportError:
            c2.info("Install python-docx to export a Word file.")
    st.caption("Budget inputs and narrative are user-supplied demo values, not verified JGVS data.")
