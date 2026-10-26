"""Organisation profile and source-document review."""
import logging

import streamlit as st

from state import org, update_org_profile, save_org_fact_review
from ui.components import fact_card, page_header
from services.export import escape_html

LOGGER = logging.getLogger(__name__)
FACTS = [
    "Legal name and registration number", "12A / 80G status and validity", "CSR-1 registration status",
    "Audited statements and annual reports", "Project locations and beneficiary counts",
    "Past outcomes with source documents", "Bank and authorised signatory details",
]


def _extract_text(upload) -> str:
    """Extract text from supported documents with a friendly fallback."""
    name = upload.name.lower()
    try:
        if name.endswith(".txt"):
            return upload.getvalue().decode("utf-8", errors="replace")
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(upload)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if name.endswith(".docx"):
            from docx import Document
            document = Document(upload)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            tables = ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
            return "\n".join([*paragraphs, *tables])
    except ImportError:
        return "Install document extraction dependencies from requirements.txt, then restart the app."
    except Exception:  # parser errors are expected for malformed uploads
        LOGGER.exception("Could not parse uploaded source document %s", upload.name)
        return "Could not extract text from this file. Review the original document directly."
    return "Unsupported document type."


def render() -> None:
    """Render editable profile fields, source uploads and verified-fact meter."""
    page_header("Knowledge Base", "Review source documents and keep organisation facts explicitly verified.")
    profile = org()
    st.subheader(profile.name)
    st.caption("Focus areas · " + " · ".join(profile.focus))
    with st.form("org_profile_form"):
        registration = st.text_input("Legal registration", profile.registration)
        tax_status = st.text_input("12A / 80G status", profile.tax_status)
        csr1 = st.text_input("CSR-1", profile.csr1)
        annual_budget = st.text_input("Annual budget", profile.annual_budget)
        beneficiaries = st.text_input("Beneficiary count", profile.beneficiaries)
        if st.form_submit_button("Save profile fields", type="primary"):
            values = {"registration": registration, "tax_status": tax_status, "csr1": csr1, "annual_budget": annual_budget, "beneficiaries": beneficiaries}
            update_org_profile(values)
            st.toast("Profile saved. Values are not verified until checked against source documents.")
            st.rerun()
    fact_cols = st.columns(3)
    for index, (label, value) in enumerate([("Registration", profile.registration), ("12A / 80G", profile.tax_status), ("CSR-1", profile.csr1), ("Annual budget", profile.annual_budget), ("Beneficiaries", profile.beneficiaries)]):
        with fact_cols[index % 3]: fact_card(label, value)
    uploads = st.file_uploader("Add source documents for review", type=["pdf", "docx", "txt"], accept_multiple_files=True)
    if uploads:
        st.subheader("Documents in this session")
        for file in uploads:
            icon = {".pdf": "PDF", ".docx": "DOCX", ".txt": "TXT"}.get("." + file.name.rsplit(".", 1)[-1].lower(), "FILE")
            with st.expander(f"{icon}  {escape_html(file.name)} · {file.size:,} bytes"):
                text = _extract_text(file)
                st.text_area("Extracted text preview", text[:12000], height=200, key="doc_preview_" + file.name)
                if "Could not parse" in text:
                    st.error(text)
                st.caption("Preview is session-only. Check the original document before marking a fact verified.")
    saved = profile.facts
    checks = profile.facts
    st.subheader("Verified facts checklist")
    done = 0
    for fact in FACTS:
        key = "fact_" + fact
        checked = st.checkbox(fact, value=bool(checks.get(fact)), key=key)
        if checked != bool(checks.get(fact)):
            save_org_fact_review(fact, checked)
        done += checked
    st.progress(done / len(FACTS), text=f"Profile completeness · {done} of {len(FACTS)} facts reviewed")
    st.caption("Demo profile only. No JGVS fact is treated as verified without human source review.")
